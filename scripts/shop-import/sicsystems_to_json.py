#!/usr/bin/env python3
"""Convert the SICSYSTEMS Google Sheet (downloaded as .xlsx) into JSON that
scripts/shop-import/import.cjs can load into the shop database.

Usage:  python3 sicsystems_to_json.py SICSYSTEMS.xlsx out.json
Needs:  pip install openpyxl

Reads every per-invoice tab (title contains a 4-digit number) by finding its
labels ("Customer Name", "Description", "Labor Subtotal", ...), so small layout
differences between tabs are fine. Also reads Jobs_Log (status, payment, our
parts cost) and Settings (tax rate, shop details).

The output holds customer details. Keep it private: do not commit it.
"""
import datetime, json, re, sys
import openpyxl

def text(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v).strip()

def num(v):
    if v is None or v == '': return None
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace('$', '').replace(',', '')
    try: return float(s)
    except ValueError: return None

def find(ws, label, max_row=60):
    want = label.lower()
    for row in ws.iter_rows(min_row=1, max_row=max_row):
        for c in row:
            if isinstance(c.value, str) and c.value.strip().lower() == want:
                return c.row, c.column
    return None

def right_of(ws, label, max_row=60):
    pos = find(ws, label, max_row)
    if not pos: return None
    v = ws.cell(pos[0], pos[1] + 1).value   # only the cell directly to the right
    return None if v in (None, '') else v

def iso(v):
    if isinstance(v, datetime.datetime): return v.date().isoformat()
    if isinstance(v, datetime.date): return v.isoformat()
    return None

def invoice(ws):
    head = text(ws['A1'].value)
    m = re.search(r'(\d{4})', head) or re.search(r'(\d{4})', ws.title)
    number = int(m.group(1)) if m else None
    hdr = find(ws, 'Description')
    end = find(ws, 'Labor Subtotal', 120)
    if not hdr or not end: return None
    cols = {text(ws.cell(hdr[0], c).value).lower(): c for c in range(1, 15) if ws.cell(hdr[0], c).value}
    def col(name): return cols.get(name.lower())
    lines = []
    for r in range(hdr[0] + 1, end[0]):
        desc = text(ws.cell(r, col('Description')).value)
        hrs = num(ws.cell(r, col('Labor Hrs')).value) if col('Labor Hrs') else None
        labor = num(ws.cell(r, col('Labor $')).value) if col('Labor $') else None
        price = num(ws.cell(r, col('Cost')).value) if col('Cost') else None
        log = text(ws.cell(r, col('OrderLog#')).value) if col('OrderLog#') else ''
        ordered = text(ws.cell(r, col('Parts Ordered')).value) if col('Parts Ordered') else ''
        line_total = num(ws.cell(r, col('Line Total $')).value) if col('Line Total $') else None
        if not desc and not hrs and not labor and not price: continue
        lines.append({'row': r, 'description': desc, 'labor_hours': hrs, 'labor_amount': labor,
                      'parts_price': price, 'order_log': log, 'parts_ordered': ordered, 'line_total': line_total})
    miles = None
    for r in range(1, hdr[0]):
        for c in range(1, 6):
            v = ws.cell(r, c).value
            if isinstance(v, str) and re.search(r'mile', v, re.I):
                mm = re.search(r'([\d,]{2,})', v)
                if mm: miles = mm.group(1).replace(',', '')
    lab_sub = num(ws.cell(end[0], end[1] + 1).value)
    return {
        'tab': ws.title, 'number': number, 'header': head,
        'date': iso(right_of(ws, 'Date')),
        'status': text(right_of(ws, 'Status')),
        'customer': {k: text(right_of(ws, l)) for k, l in
                     (('name', 'Customer Name'), ('phone', 'Phone'), ('email', 'Email'), ('address', 'Address'))},
        'bike': {'year': text(right_of(ws, 'Year')), 'make': text(right_of(ws, 'Make')),
                 'model': text(right_of(ws, 'Model')), 'vin': text(right_of(ws, 'VIN')), 'mileage': miles or ''},
        'lines': lines,
        'sheet_totals': {'labor': lab_sub, 'parts': num(right_of(ws, 'Parts Subtotal', 120)),
                         'subtotal': num(right_of(ws, 'Total', 120)), 'tax': num(right_of(ws, 'Tax on Parts', 120)),
                         'grand': num(right_of(ws, 'Grand Total', 120))},
        'payment': '',
        'authorized_by': text(right_of(ws, 'Authorized By', 120)),
    }

def payment_row(ws):
    # The lower "Status" cell (Paid/Unpaid) sits next to "Authorized By".
    auth = find(ws, 'Authorized By', 120)
    if not auth: return '', ''
    r = auth[0] - 1
    return text(ws.cell(r, auth[1] + 1).value), text(ws.cell(r, auth[1] + 2).value)

def main(src, out):
    wb = openpyxl.load_workbook(src, data_only=True)
    invoices = []
    for ws in wb.worksheets:
        if not re.search(r'\d{4}', ws.title) or ws.title.lower().startswith('copy of'): continue
        inv = invoice(ws)
        if not inv: continue
        inv['payment'], inv['payment_note'] = payment_row(ws)
        invoices.append(inv)
    log = {}
    if 'Jobs_Log' in wb.sheetnames:
        ws = wb['Jobs_Log']
        heads = [text(c.value) for c in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            rec = dict(zip(heads, row))
            n = num(rec.get('Invoice #'))
            if n and (rec.get('Customer') or rec.get('Total $') is not None):
                log[int(n)] = {k: (iso(v) or v) for k, v in rec.items() if k}
    settings = {}
    if 'Settings' in wb.sheetnames:
        for row in wb['Settings'].iter_rows(min_row=1, values_only=True):
            if row and row[0] and len(row) > 1: settings[text(row[0])] = row[1]
    policies = []
    if invoices:
        ws = wb[invoices[0]['tab']]
        start = find(ws, 'Service Policy + Authorization', 120)
        if start:
            for r in range(start[0] + 1, start[0] + 15):
                v = text(ws.cell(r, 1).value)
                if v: policies.append(v.lstrip('✚ ').strip())
    json.dump({'source': src.split('/')[-1], 'settings': settings, 'policies': policies,
               'jobs_log': log, 'invoices': invoices}, open(out, 'w'), indent=1, default=str)
    print(f'{len(invoices)} invoice tabs, {len(log)} Jobs_Log rows -> {out}')

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
